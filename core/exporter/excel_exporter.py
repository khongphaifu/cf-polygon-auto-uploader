import os
from datetime import datetime
from pathlib import Path
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core.models import PolygonProblemRecord


class ExcelExporter:
    """Exports processed problem records to a beautifully styled Excel workbook."""

    HEADERS = [
        "STT",
        "Tên Bài (Title)",
        "Link Gốc (Source URL)",
        "Polygon Problem ID",
        "Link Polygon",
        "Link Add Codeforces (Import URL)",
        "Trạng Thái Package",
        "Package ID",
        "Sẵn Sàng Import CF",
        "Thời Gian Tạo",
        "Ghi Chú / Lỗi",
    ]

    def __init__(self, output_path: str = "problems_summary.xlsx"):
        self.output_path = Path(output_path)

    def _setup_headers(self, ws) -> None:
        """Create header row with professional styling."""
        ws.append(self.HEADERS)

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(self.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        ws.row_dimensions[1].height = 28

    def _apply_styles_and_widths(self, ws) -> None:
        """Auto-adjust column widths and borders."""
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(self.HEADERS)):
            for cell in row:
                cell.border = thin_border
                if cell.column in [1, 4, 7, 8, 9, 10]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center")

        # Color-code status
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        green_font = Font(color="375623", bold=True)
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        red_font = Font(color="C65911", bold=True)

        for row_idx in range(2, ws.max_row + 1):
            pkg_cell = ws.cell(row=row_idx, column=7)
            if pkg_cell.value in ["READY", "SUCCESS"]:
                pkg_cell.fill = green_fill
                pkg_cell.font = green_font
            elif pkg_cell.value in ["FAILED", "ERROR"]:
                pkg_cell.fill = red_fill
                pkg_cell.font = red_font

            cf_cell = ws.cell(row=row_idx, column=9)
            if cf_cell.value == "YES":
                cf_cell.fill = green_fill
                cf_cell.font = green_font

        # Adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 36
        ws.column_dimensions["E"].width = 42
        ws.column_dimensions["F"].width = 46
        ws.column_dimensions["K"].width = 30

    def export_records(self, records: List[PolygonProblemRecord]) -> str:
        """Export or append problem records into Excel file."""
        if self.output_path.exists():
            try:
                wb = openpyxl.load_workbook(self.output_path)
                ws = wb.active
                # Verify header match
                if [cell.value for cell in ws[1]] != self.HEADERS:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Problems"
                    self._setup_headers(ws)
            except Exception:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Problems"
                self._setup_headers(ws)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Problems"
            self._setup_headers(ws)

        # Existing IDs to avoid duplicates
        existing_ids = set()
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=4).value
            if val is not None:
                try:
                    existing_ids.add(int(val))
                except Exception:
                    pass

        stt = ws.max_row

        for rec in records:
            cf_import_url = rec.cf_import_url or (f"https://polygon.codeforces.com/p{rec.id}/{rec.owner}/{rec.name}" if rec.owner else f"https://polygon.codeforces.com/p{rec.id}/{rec.name}")
            polygon_url = rec.url or f"https://polygon.codeforces.com/p{rec.id}/{rec.name}"

            if rec.id in existing_ids:
                # Update existing row
                for r in range(2, ws.max_row + 1):
                    if ws.cell(row=r, column=4).value == rec.id:
                        ws.cell(row=r, column=6, value=cf_import_url)
                        ws.cell(row=r, column=7, value=rec.package_state or rec.status)
                        ws.cell(row=r, column=8, value=rec.package_id)
                        ws.cell(row=r, column=9, value="YES" if rec.package_state == "READY" else "PENDING")
                        ws.cell(row=r, column=11, value=rec.error_message or "")
                        break
            else:
                stt += 1
                created_time = rec.created_at or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                is_cf_ready = "YES" if (rec.package_state == "READY" or rec.status == "SUCCESS") else "NO"

                row_data = [
                    stt - 1,
                    rec.name,
                    rec.source_url,
                    rec.id,
                    polygon_url,
                    cf_import_url,
                    rec.package_state or rec.status,
                    rec.package_id or "-",
                    is_cf_ready,
                    created_time,
                    rec.error_message or "",
                ]
                ws.append(row_data)

                # Set hyperlinks
                current_row = ws.max_row
                if rec.source_url:
                    ws.cell(row=current_row, column=3).hyperlink = rec.source_url
                    ws.cell(row=current_row, column=3).font = Font(color="0563C1", underline="single")
                if polygon_url:
                    ws.cell(row=current_row, column=5).hyperlink = polygon_url
                    ws.cell(row=current_row, column=5).font = Font(color="0563C1", underline="single")
                if cf_import_url:
                    ws.cell(row=current_row, column=6).hyperlink = cf_import_url
                    ws.cell(row=current_row, column=6).font = Font(color="0563C1", underline="single")

        self._apply_styles_and_widths(ws)
        wb.save(self.output_path)
        return str(self.output_path.resolve())
